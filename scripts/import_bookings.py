"""Data migration: load bookings from a customer's spreadsheet export (CSV or Excel) into the database.

    python -m scripts.import_bookings bookings.xlsx [--dry-run]

Expected columns (case-insensitive): reference, guest_name, room_type, check_in, check_out, status.
Every row is validated; bad rows are reported with their row number and skipped, and existing references
are updated rather than duplicated, so the import is safe to re-run.
"""

from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Booking
from app.tools import BOOKING_REF

REQUIRED = ["reference", "guest_name", "room_type", "check_in", "check_out"]
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y")


@dataclass
class ImportReport:
    inserted: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)


def read_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(h).strip().lower() for h in rows[0]]
        return [{header[i]: ("" if v is None else v) for i, v in enumerate(r)} for r in rows[1:]]
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [{k.strip().lower(): v for k, v in row.items()} for row in csv.DictReader(f)]


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date '{value}'")


def import_bookings(session: Session, rows: list[dict], dry_run: bool = False) -> ImportReport:
    report = ImportReport()
    seen: set[str] = set()
    for line, row in enumerate(rows, start=2):  # row 1 is the header
        try:
            missing = [c for c in REQUIRED if not str(row.get(c, "")).strip()]
            if missing:
                raise ValueError(f"missing {', '.join(missing)}")
            reference = str(row["reference"]).strip().upper()
            if not BOOKING_REF.fullmatch(reference):
                raise ValueError(f"bad reference '{reference}'")
            if reference in seen:
                raise ValueError(f"duplicate reference '{reference}' in file")
            check_in, check_out = parse_date(row["check_in"]), parse_date(row["check_out"])
            if check_out <= check_in:
                raise ValueError("check_out must be after check_in")
        except ValueError as exc:
            report.errors.append(f"row {line}: {exc}")
            continue

        seen.add(reference)
        values = {
            "guest_name": str(row["guest_name"]).strip(),
            "room_type": str(row["room_type"]).strip(),
            "check_in": check_in,
            "check_out": check_out,
            "status": str(row.get("status") or "confirmed").strip().lower(),
        }
        existing = session.scalar(select(Booking).where(Booking.reference == reference))
        if existing:
            for key, value in values.items():
                setattr(existing, key, value)
            report.updated += 1
        else:
            session.add(Booking(reference=reference, **values))
            report.inserted += 1

    if dry_run:
        session.rollback()
    else:
        session.commit()
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    from app.db import SessionLocal, init_db

    init_db()
    with SessionLocal() as session:
        report = import_bookings(session, read_rows(args.path), args.dry_run)
    print(f"inserted={report.inserted} updated={report.updated} errors={len(report.errors)}")
    for err in report.errors:
        print("  " + err)
    return 1 if report.errors else 0


if __name__ == "__main__":
    sys.exit(main())
